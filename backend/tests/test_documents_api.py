import io
import logging
import pathlib
import tempfile
import uuid
from datetime import datetime, timedelta, timezone

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import sessionmaker

from app import database
from app.database import Base, get_db
from app.main import app
from app.models.document import Document
from app.schemas.classification import ClassificationResult
from app.services import classification_service, file_service

DOCUMENT_TYPES, INSTITUTIONS = classification_service.load_catalogs()
DOCUMENT_TYPE = DOCUMENT_TYPES[0]["id"]
DOCUMENT_TYPE_NAME = DOCUMENT_TYPES[0]["name"]
INSTITUTION = INSTITUTIONS[0]["id"]
INSTITUTION_NAME = INSTITUTIONS[0]["name"]
CLASSIFY_URL = "/api/documents/classify"
SECRET_MARKER = "GIZLI_BELGE_ICERIGI"  # loglarda görünmemesi gereken belge metni işareti
PDF_TEXT = f"Sayin yetkili, sokagimizdaki copler toplanmiyor. {SECRET_MARKER}"
DOCX_TEXT = f"Sayın yetkili, parktaki salıncak kırık. {SECRET_MARKER}"
RESPONSE_FIELDS = {
    "document_id", "file_name", "file_type", "document_type", "document_type_name", "institution_id", "institution_name",
    "needs_review", "review_reason", "summary", "sender_name", "sender_institution", "status",
}
SUMMARY = "Vatandaş sokaktaki çöplerin toplanmadığını bildirip gereğinin yapılmasını istiyor."
SENDER_NAME = "Ayşe Yılmaz"
SENDER_INSTITUTION = "Çiğdem Mahallesi Muhtarlığı"
TEXT_FAILED_MESSAGE = "Belgeden sınıflandırma için yeterli metin çıkarılamadı."
CLASSIFICATION_FAILED_MESSAGE = "Belge şu anda sınıflandırılamadı. Lütfen daha sonra tekrar deneyin."


def make_pdf(text: str) -> bytes:
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
        return pdf.tobytes()


def make_docx(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def classification(needs_review: bool = False, sender: bool = True) -> ClassificationResult:
    """Sahte sınıflandırma sonucu. sender=False: belgede gönderen bilgisi yok (D-044)."""
    extra = {
        "summary": SUMMARY,
        "sender_name": SENDER_NAME if sender else None,
        "sender_institution": SENDER_INSTITUTION if sender else None,
    }
    if needs_review:
        return ClassificationResult(document_type="other", institution_id=None, needs_review=True, review_reason="Kurum belirsiz.", **extra)
    return ClassificationResult(document_type=DOCUMENT_TYPE, institution_id=INSTITUTION, needs_review=False, review_reason=None, **extra)


def classification_error() -> classification_service.ClassificationError:
    error = classification_service.ClassificationError("Belge Gemini ile sınıflandırılamadı.")
    error.__cause__ = RuntimeError("ham Gemini hata detayı: API_KEY_INVALID")
    return error


def upload(client: TestClient, file_name: str, content: bytes):
    return client.post(CLASSIFY_URL, files={"file": (file_name, content, "application/octet-stream")})


def all_documents(session_factory) -> list[Document]:
    with session_factory() as session:
        return list(session.scalars(select(Document)))


def stored_files(storage_dir) -> list[str]:
    return sorted(path.name for path in storage_dir.iterdir()) if storage_dir.exists() else []


@pytest.fixture(autouse=True)
def gemini_must_not_be_called_unexpectedly(monkeypatch):
    def unexpected_classify_text(text):
        raise AssertionError("classify_text beklenmedik şekilde çağrıldı")

    monkeypatch.setattr(classification_service, "classify_text", unexpected_classify_text)


@pytest.fixture
def fake_classify(monkeypatch):
    """classification_service.classify_text'i sahte bir sonuç veya exception ile değiştirir; gelen metinleri kaydeder."""
    calls = []

    def install(outcome):
        def classify_text(text):
            calls.append(text)
            if isinstance(outcome, Exception):
                raise outcome
            return outcome

        monkeypatch.setattr(classification_service, "classify_text", classify_text)
        return calls

    return install


@pytest.fixture
def storage_dir(tmp_path, monkeypatch):
    path = tmp_path / "storage"
    monkeypatch.setattr(file_service, "STORAGE_DIR", path)
    return path


@pytest.fixture
def session_factory(tmp_path):
    # İzole SQLite test veritabanı. create_all yalnızca testte kullanılır; gerçek şema Alembic ile yönetilir (D-030).
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    yield sessionmaker(bind=engine)
    engine.dispose()


@pytest.fixture
def client(session_factory, storage_dir):
    def override_get_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    yield TestClient(app, raise_server_exceptions=False)
    app.dependency_overrides.clear()


# --- Başarılı akış ---


def test_health_still_works(client):
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


@pytest.mark.parametrize(
    "file_name, content, file_type, text",
    [("dilekce.pdf", make_pdf(PDF_TEXT), "pdf", PDF_TEXT), ("dilekce.docx", make_docx(DOCX_TEXT), "docx", DOCX_TEXT)],
    ids=["pdf", "docx"],
)
def test_valid_document_is_classified(client, session_factory, storage_dir, fake_classify, caplog, file_name, content, file_type, text):
    calls = fake_classify(classification())

    with caplog.at_level(logging.DEBUG):
        response = upload(client, file_name, content)

    assert response.status_code == 200
    body = response.json()
    assert set(body) == RESPONSE_FIELDS  # file_reference ve extracted_text yanıtta yok
    assert body == {
        "document_id": body["document_id"], "file_name": file_name, "file_type": file_type,
        "document_type": DOCUMENT_TYPE, "document_type_name": DOCUMENT_TYPE_NAME,
        "institution_id": INSTITUTION, "institution_name": INSTITUTION_NAME,
        "needs_review": False, "review_reason": None, "status": "classified",
        "summary": SUMMARY, "sender_name": SENDER_NAME, "sender_institution": SENDER_INSTITUTION,
    }
    [document] = all_documents(session_factory)
    assert str(document.id) == body["document_id"]
    assert document.extracted_text == text
    assert calls == [text]
    assert document.file_reference == f"{document.id}.{file_type}"
    assert (storage_dir / document.file_reference).read_bytes() == content
    assert SECRET_MARKER not in caplog.text


def test_needs_review_result_returns_200_with_needs_review_status(client, session_factory, fake_classify):
    fake_classify(classification(needs_review=True))

    response = upload(client, "dilekce.docx", make_docx(DOCX_TEXT))

    assert response.status_code == 200
    body = response.json()
    assert (body["status"], body["needs_review"], body["institution_id"], body["review_reason"]) == (
        "needs_review", True, None, "Kurum belirsiz.",
    )
    # Kurum atanmadığında ad da null; belge türü adı yine katalogdan gelir.
    assert (body["institution_name"], body["document_type_name"]) == (None, classification_service.DOCUMENT_TYPE_NAMES["other"])
    [document] = all_documents(session_factory)
    assert (document.status, document.needs_review, document.institution_id, document.review_reason) == (
        "needs_review", True, None, "Kurum belirsiz.",
    )


def test_full_extracted_text_is_stored_and_passed_to_classification(client, session_factory, fake_classify):
    long_text = "a" * 60_000
    calls = fake_classify(classification())

    response = upload(client, "uzun.docx", make_docx(long_text))

    assert response.status_code == 200
    [document] = all_documents(session_factory)
    assert document.extracted_text == long_text
    assert calls == [long_text]  # 50.000 karakter kesmesi classification_service içinde yapılır


# --- Kabul öncesi red: kayıt yok, dosya yok ---


def test_file_over_size_limit_returns_413_without_db_or_storage(client, session_factory, storage_dir, monkeypatch):
    content = make_pdf(PDF_TEXT)
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", len(content) - 1)  # 50 MB dosya üretmeden aynı sınır mantığı

    response = upload(client, "buyuk.pdf", content)

    assert response.status_code == 413
    assert response.json() == {"detail": "Dosya boyutu 50 MB sınırını aşıyor."}
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_file_exactly_at_size_limit_is_accepted(client, session_factory, monkeypatch, fake_classify):
    content = make_pdf(PDF_TEXT)
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", len(content))
    fake_classify(classification())

    assert upload(client, "sinirda.pdf", content).status_code == 200
    assert len(all_documents(session_factory)) == 1


def test_upload_is_read_with_size_limit_not_fully(client, monkeypatch):
    read_sizes = []
    original_read = tempfile.SpooledTemporaryFile.read

    def recording_read(self, *args):
        read_sizes.append(args[0] if args else None)
        return original_read(self, *args)

    monkeypatch.setattr(tempfile.SpooledTemporaryFile, "read", recording_read)
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", 1000)

    response = upload(client, "buyuk.pdf", b"%PDF" + b"0" * 5000)

    assert response.status_code == 413
    assert read_sizes == [1001]


@pytest.mark.parametrize(
    "file_name, content",
    [
        ("dilekce.txt", b"duz metin dosyasi"),
        ("dilekce.doc", make_docx(DOCX_TEXT)),
        ("dilekce.pdf", b"PDF olmayan icerik"),
        ("dilekce.docx", b"PK ama docx degil"),
    ],
    ids=["txt", "doc", "fake-pdf", "fake-docx"],
)
def test_unsupported_file_returns_415_without_db_or_storage(client, session_factory, storage_dir, file_name, content):
    response = upload(client, file_name, content)

    assert response.status_code == 415
    assert response.json() == {"detail": "Yalnızca PDF, DOC, DOCX, JPG, JPEG veya PNG dosyaları kabul edilir."}
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_missing_file_returns_validation_error_without_db_or_storage(client, session_factory, storage_dir):
    response = client.post(CLASSIFY_URL)

    assert response.status_code == 422
    assert "document_id" not in response.json()
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


# --- Kabul sonrası failed kayıtları ---


@pytest.mark.parametrize(
    "file_name, content, expected_text",
    [
        ("kisa.docx", make_docx("Kısa"), "Kısa"),
        ("metinsiz.pdf", make_pdf(""), None),
        ("bozuk.pdf", b"%PDF-1.7\nbozuk icerik", None),
    ],
    ids=["short-text", "no-text", "corrupt"],
)
def test_text_extraction_failure_returns_422_with_failed_record(client, session_factory, storage_dir, file_name, content, expected_text):
    response = upload(client, file_name, content)

    assert response.status_code == 422
    body = response.json()
    assert set(body) == RESPONSE_FIELDS | {"message"}
    assert body["message"] == TEXT_FAILED_MESSAGE
    assert (
        body["status"], body["document_type"], body["document_type_name"], body["institution_id"], body["institution_name"],
        body["needs_review"], body["review_reason"],
    ) == ("failed", None, None, None, None, False, None)
    [document] = all_documents(session_factory)
    assert str(document.id) == body["document_id"]
    assert (document.status, document.document_type, document.institution_id, document.needs_review, document.review_reason) == (
        "failed", None, None, False, None,
    )
    assert document.extracted_text == expected_text
    assert (storage_dir / document.file_reference).read_bytes() == content


def test_classification_error_returns_502_with_failed_record(client, session_factory, storage_dir, fake_classify, caplog):
    fake_classify(classification_error())
    content = make_docx(DOCX_TEXT)

    with caplog.at_level(logging.DEBUG):
        response = upload(client, "dilekce.docx", content)

    assert response.status_code == 502
    body = response.json()
    assert set(body) == RESPONSE_FIELDS | {"message"}
    assert body["message"] == CLASSIFICATION_FAILED_MESSAGE
    assert "ham Gemini" not in response.text and "API_KEY_INVALID" not in response.text
    assert (
        body["status"], body["document_type"], body["document_type_name"], body["institution_id"], body["institution_name"],
        body["needs_review"], body["review_reason"],
    ) == ("failed", None, None, None, None, False, None)
    [document] = all_documents(session_factory)
    assert (document.status, document.document_type, document.institution_id, document.needs_review, document.review_reason) == (
        "failed", None, None, False, None,
    )
    assert document.extracted_text == DOCX_TEXT
    assert (storage_dir / document.file_reference).read_bytes() == content
    assert SECRET_MARKER not in caplog.text


# --- Tutarlılık ---


def test_response_names_are_read_from_catalogs_and_not_stored(client, session_factory, fake_classify, monkeypatch):
    fake_classify(classification())
    monkeypatch.setitem(classification_service.DOCUMENT_TYPE_NAMES, DOCUMENT_TYPE, "Katalogdan Gelen Tür")
    monkeypatch.setitem(classification_service.INSTITUTION_NAMES, INSTITUTION, "Katalogdan Gelen Kurum")

    body = upload(client, "dilekce.docx", make_docx(DOCX_TEXT)).json()

    assert (body["document_type_name"], body["institution_name"]) == ("Katalogdan Gelen Tür", "Katalogdan Gelen Kurum")
    assert (body["document_type"], body["institution_id"]) == (DOCUMENT_TYPE, INSTITUTION)  # ID davranışı değişmez
    [document] = all_documents(session_factory)
    assert not hasattr(document, "document_type_name") and not hasattr(document, "institution_name")  # adlar DB'ye yazılmaz


def test_user_file_name_is_not_used_as_storage_path(client, session_factory, storage_dir, fake_classify, tmp_path):
    fake_classify(classification())
    user_file_name = "../../gizli/dilekce.pdf"

    response = upload(client, user_file_name, make_pdf(PDF_TEXT))

    assert response.status_code == 200
    body = response.json()
    document_id = uuid.UUID(body["document_id"])
    [document] = all_documents(session_factory)
    assert document.id == document_id
    assert document.file_name == body["file_name"] == user_file_name
    assert "gizli" not in document.file_reference and document.file_reference == f"{document_id}.pdf"
    assert stored_files(storage_dir) == [f"{document_id}.pdf"]
    assert list(tmp_path.rglob("*.pdf")) == [storage_dir / f"{document_id}.pdf"]


@pytest.mark.parametrize(
    "content, outcome",
    [
        (make_docx(DOCX_TEXT), classification()),
        (make_docx("Kısa"), None),
        (make_docx(DOCX_TEXT), classification_error()),
    ],
    ids=["classified", "text-failed", "classification-failed"],
)
def test_db_commit_failure_rolls_back_and_removes_orphan_file(client, session_factory, storage_dir, fake_classify, caplog, content, outcome):
    rollbacks = []

    def failing_get_db():
        with session_factory() as session:
            original_rollback = session.rollback

            def fail_commit():
                raise OperationalError("COMMIT", None, Exception("veritabanı erişilemiyor"))

            def tracking_rollback():
                rollbacks.append(True)
                original_rollback()

            session.commit = fail_commit
            session.rollback = tracking_rollback
            yield session

    app.dependency_overrides[get_db] = failing_get_db
    if outcome is not None:
        fake_classify(outcome)

    with caplog.at_level(logging.DEBUG):
        response = upload(client, "dilekce.docx", content)

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
    assert rollbacks == [True]
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []
    assert SECRET_MARKER not in caplog.text


def test_storage_write_failure_returns_500_without_partial_file_or_record(
    client, session_factory, storage_dir, failing_storage_write, caplog
):
    with caplog.at_level(logging.DEBUG):
        response = upload(client, "dilekce.pdf", make_pdf(PDF_TEXT))

    assert response.status_code == 500
    assert response.text == "Internal Server Error"  # "No space left" gibi teknik ayrıntı yanıta girmez
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []
    assert SECRET_MARKER not in caplog.text


def test_openapi_documents_both_422_bodies():
    responses = app.openapi()["paths"][CLASSIFY_URL]["post"]["responses"]

    schema_422 = responses["422"]["content"]["application/json"]["schema"]
    assert {item["$ref"].rsplit("/", 1)[-1] for item in schema_422["anyOf"]} == {
        "FailedClassifyResponse", "ValidationErrorResponse",
    }
    assert responses["502"]["content"]["application/json"]["schema"]["$ref"].endswith("/FailedClassifyResponse")


def test_production_engine_hides_sql_parameters():
    assert database.engine.hide_parameters is True


# --- Kayıt görünürlüğü endpoint'leri (V1.2, D-043) ---

LIST_URL = "/api/documents"


def insert_document(session_factory, **overrides) -> Document:
    """Testte doğrudan kayıt ekler; created_at sırasını deterministik kurmak için kullanılır."""
    document_id = overrides.pop("id", uuid.uuid4())
    fields = {
        "id": document_id,
        "file_name": "belge.pdf",
        "file_type": "pdf",
        "file_reference": f"{document_id}.pdf",
        "extracted_text": "Bu belgenin çıkarılmış metni.",
        "document_type": DOCUMENT_TYPE,
        "institution_id": INSTITUTION,
        "needs_review": False,
        "review_reason": None,
        "status": "classified",
        "created_at": datetime.now(timezone.utc),
    }
    fields.update(overrides)
    with session_factory() as session:
        document = Document(**fields)
        session.add(document)
        session.commit()
        session.refresh(document)
        session.expunge(document)
    return document


def test_list_returns_documents_newest_first(client, session_factory):
    now = datetime.now(timezone.utc)
    insert_document(session_factory, file_name="orta.pdf", created_at=now - timedelta(hours=1))
    insert_document(session_factory, file_name="en_eski.pdf", created_at=now - timedelta(hours=2))
    insert_document(session_factory, file_name="en_yeni.pdf", created_at=now)

    response = client.get(LIST_URL)

    assert response.status_code == 200
    assert [item["file_name"] for item in response.json()] == ["en_yeni.pdf", "orta.pdf", "en_eski.pdf"]


def test_list_is_empty_when_no_documents(client):
    response = client.get(LIST_URL)
    assert response.status_code == 200
    assert response.json() == []


def test_list_item_has_catalog_names_and_created_at(client, session_factory):
    insert_document(session_factory)

    item = client.get(LIST_URL).json()[0]

    assert item["document_type"] == DOCUMENT_TYPE
    assert item["document_type_name"] == DOCUMENT_TYPE_NAME
    assert item["institution_id"] == INSTITUTION
    assert item["institution_name"] == INSTITUTION_NAME
    assert set(item) == RESPONSE_FIELDS | {"created_at"}


def test_list_hides_extracted_text_and_file_reference(client, session_factory):
    document = insert_document(session_factory, extracted_text=SECRET_MARKER)

    body = client.get(LIST_URL).text

    assert "extracted_text" not in body
    assert "file_reference" not in body
    assert SECRET_MARKER not in body
    assert document.file_reference not in body


def test_needs_review_item_keeps_review_reason_and_null_institution(client, session_factory):
    insert_document(
        session_factory, status="needs_review", needs_review=True,
        institution_id=None, review_reason="Kurum belirsiz.",
    )

    item = client.get(LIST_URL).json()[0]

    assert item["status"] == "needs_review"
    assert item["needs_review"] is True
    assert item["institution_id"] is None and item["institution_name"] is None
    assert item["review_reason"] == "Kurum belirsiz."


def test_detail_returns_extracted_text_but_not_file_reference(client, session_factory):
    document = insert_document(session_factory, extracted_text=SECRET_MARKER)

    response = client.get(f"{LIST_URL}/{document.id}")

    assert response.status_code == 200
    assert response.json()["extracted_text"] == SECRET_MARKER
    assert set(response.json()) == RESPONSE_FIELDS | {"created_at", "extracted_text"}
    assert "file_reference" not in response.text
    assert document.file_reference not in response.text


def test_detail_returns_null_extracted_text_for_failed_document(client, session_factory):
    document = insert_document(
        session_factory, status="failed", extracted_text=None,
        document_type=None, institution_id=None,
    )

    body = client.get(f"{LIST_URL}/{document.id}").json()

    assert body["extracted_text"] is None
    assert body["status"] == "failed"
    assert body["document_type_name"] is None and body["institution_name"] is None


def test_detail_returns_404_for_unknown_document(client):
    response = client.get(f"{LIST_URL}/{uuid.uuid4()}")

    assert response.status_code == 404
    assert response.json() == {"detail": "Belge bulunamadı."}


def test_download_returns_original_pdf_bytes_with_original_file_name(client, fake_classify, storage_dir):
    fake_classify(classification())
    content = make_pdf(PDF_TEXT)
    document_id = upload(client, "dilekçe raporu.pdf", content).json()["document_id"]

    response = client.get(f"{LIST_URL}/{document_id}/download")

    assert response.status_code == 200
    assert response.content == content  # byte-for-byte aynı dosya
    assert response.headers["content-type"] == "application/pdf"
    # Türkçe/boşluklu ad RFC 5987 ile kodlanır; storage adı (UUID) sızmaz.
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert "dilek" in disposition
    assert document_id not in disposition


def test_download_returns_docx_with_correct_media_type(client, fake_classify):
    fake_classify(classification())
    content = make_docx(DOCX_TEXT)
    document_id = upload(client, "basvuru.docx", content).json()["document_id"]

    response = client.get(f"{LIST_URL}/{document_id}/download")

    assert response.status_code == 200
    assert response.content == content
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert 'filename="basvuru.docx"' in response.headers["content-disposition"]


def test_download_returns_404_for_unknown_document(client):
    response = client.get(f"{LIST_URL}/{uuid.uuid4()}/download")

    assert response.status_code == 404
    assert response.json() == {"detail": "Belge bulunamadı."}


def test_download_returns_404_without_details_when_stored_file_is_missing(client, session_factory, storage_dir, caplog):
    document = insert_document(session_factory)  # kayıt var, storage dosyası hiç yazılmadı

    with caplog.at_level(logging.WARNING):
        response = client.get(f"{LIST_URL}/{document.id}/download")

    assert response.status_code == 404
    assert response.json() == {"detail": "Belge bulunamadı."}
    assert document.file_reference not in response.text
    assert str(storage_dir) not in response.text
    assert document.file_reference not in caplog.text  # log da dosya adını yazmaz


# --- Özet ve gönderen bilgisi (V1.2 · Adım 2, D-044) ---


def test_classified_document_stores_and_returns_summary_and_sender(client, fake_classify, session_factory):
    fake_classify(classification())

    body = upload(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()

    assert body["summary"] == SUMMARY
    assert body["sender_name"] == SENDER_NAME
    assert body["sender_institution"] == SENDER_INSTITUTION
    [document] = all_documents(session_factory)
    assert (document.summary, document.sender_name, document.sender_institution) == (
        SUMMARY, SENDER_NAME, SENDER_INSTITUTION,
    )


def test_document_without_sender_information_keeps_sender_fields_null(client, fake_classify, session_factory):
    fake_classify(classification(sender=False))

    body = upload(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()

    assert body["summary"] == SUMMARY  # özet yine dolu
    assert body["sender_name"] is None
    assert body["sender_institution"] is None
    [document] = all_documents(session_factory)
    assert document.sender_name is None and document.sender_institution is None


def test_needs_review_document_keeps_summary(client, fake_classify, session_factory):
    fake_classify(classification(needs_review=True))

    body = upload(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()

    assert body["status"] == "needs_review"
    assert body["summary"] == SUMMARY
    [document] = all_documents(session_factory)
    assert document.summary == SUMMARY


@pytest.mark.parametrize(
    "file_name, content, outcome",
    [
        ("kisa.pdf", None, None),  # metin çıkarılamadı → 422
        ("dilekce.pdf", None, "error"),  # Gemini hatası → 502
    ],
)
def test_failed_document_leaves_summary_and_sender_null(
    client, fake_classify, session_factory, file_name, content, outcome
):
    if outcome == "error":
        fake_classify(classification_error())
        payload = make_pdf(PDF_TEXT)
    else:
        payload = make_pdf("kisa")  # 10 karakterin altında → Gemini'ye gitmez

    body = upload(client, file_name, payload).json()

    assert body["status"] == "failed"
    assert body["summary"] is None
    assert body["sender_name"] is None
    assert body["sender_institution"] is None
    [document] = all_documents(session_factory)
    assert (document.summary, document.sender_name, document.sender_institution) == (None, None, None)


def test_list_and_detail_return_summary_and_sender(client, fake_classify):
    fake_classify(classification())
    document_id = upload(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()["document_id"]

    [item] = client.get(LIST_URL).json()
    detail = client.get(f"{LIST_URL}/{document_id}").json()

    for body in (item, detail):
        assert body["summary"] == SUMMARY
        assert body["sender_name"] == SENDER_NAME
        assert body["sender_institution"] == SENDER_INSTITUTION
    assert "file_reference" not in client.get(LIST_URL).text


def test_old_record_without_summary_is_returned_with_null_fields(client, session_factory):
    """Migration öncesi yazılmış kayıtlarda üç alan da null'dır; liste ve detay bunu olduğu gibi döner."""
    document = insert_document(session_factory, summary=None, sender_name=None, sender_institution=None)

    [item] = client.get(LIST_URL).json()
    detail = client.get(f"{LIST_URL}/{document.id}").json()

    for body in (item, detail):
        assert body["summary"] is None
        assert body["sender_name"] is None
        assert body["sender_institution"] is None


def test_model_columns_match_migrated_schema():
    """Model ile Alembic şeması aynı kolonları taşımalı (D-030); alembic check bunu ayrıca doğrular."""
    columns = set(Document.__table__.columns.keys())
    assert {"summary", "sender_name", "sender_institution"} <= columns
    for name in ("summary", "sender_name", "sender_institution"):
        assert Document.__table__.columns[name].nullable is True


# --- JPG / JPEG / PNG desteği ---

IMAGE_TEXT = f"Sayin yetkili, sokaktaki cukur onarilmali. {SECRET_MARKER}"


def make_image(text: str = IMAGE_TEXT, image_format: str = "jpeg") -> bytes:
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
        pixmap = page.get_pixmap(dpi=96)
        return pixmap.tobytes("jpeg", jpg_quality=90) if image_format == "jpeg" else pixmap.tobytes("png")


@pytest.fixture
def fake_image_ocr(monkeypatch):
    """Görüntü OCR'ını sabit metinle değiştirir; testler gerçek Tesseract gerektirmez."""
    def install(text=IMAGE_TEXT):
        monkeypatch.setattr(file_service, "_ocr_page_text", lambda page: text)

    return install


@pytest.mark.parametrize(
    "file_name, image_format, expected_type",
    [("foto.jpg", "jpeg", "jpg"), ("foto.jpeg", "jpeg", "jpeg"), ("foto.png", "png", "png")],
)
def test_image_upload_is_classified(client, fake_classify, fake_image_ocr, session_factory, storage_dir,
                                    file_name, image_format, expected_type):
    fake_image_ocr()
    calls = fake_classify(classification())

    response = upload(client, file_name, make_image(image_format=image_format))

    assert response.status_code == 200
    body = response.json()
    assert set(body) == RESPONSE_FIELDS
    assert body["file_type"] == expected_type
    assert body["file_name"] == file_name
    assert body["status"] == "classified"
    assert body["institution_id"] == INSTITUTION
    assert calls == [IMAGE_TEXT]  # OCR metni tek Gemini çağrısına gitti
    document = all_documents(session_factory)[0]
    assert document.extracted_text == IMAGE_TEXT
    assert stored_files(storage_dir) == [f"{body['document_id']}.{expected_type}"]


def test_image_without_enough_text_is_failed_with_422(client, fake_image_ocr, session_factory):
    fake_image_ocr("kısa")

    response = upload(client, "foto.png", make_image(image_format="png"))

    assert response.status_code == 422
    body = response.json()
    assert body["status"] == "failed"
    assert body["message"] == TEXT_FAILED_MESSAGE
    assert body["document_type"] is None and body["institution_id"] is None
    assert all_documents(session_factory)[0].status == "failed"


@pytest.mark.parametrize(
    "file_name, content",
    [
        ("sahte.jpg", b"duz metin dosyasi"),
        ("sahte.png", b"duz metin dosyasi"),
        ("foto.gif", make_image(image_format="png")),
    ],
    ids=["fake-jpg", "fake-png", "gif"],
)
def test_unsupported_image_returns_415(client, session_factory, storage_dir, file_name, content):
    response = upload(client, file_name, content)

    assert response.status_code == 415
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_oversized_image_returns_413(client, session_factory, storage_dir, monkeypatch):
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", 1000)

    response = upload(client, "foto.jpg", make_image() + b"\x00" * 1001)

    assert response.status_code == 413
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


@pytest.mark.parametrize(
    "file_name, image_format, media_type",
    [
        ("foto.jpg", "jpeg", "image/jpeg"),
        ("foto.jpeg", "jpeg", "image/jpeg"),
        ("tarama raporu.png", "png", "image/png"),
    ],
)
def test_image_download_returns_original_bytes_and_media_type(client, fake_classify, fake_image_ocr,
                                                              file_name, image_format, media_type):
    fake_image_ocr()
    fake_classify(classification())
    content = make_image(image_format=image_format)
    document_id = upload(client, file_name, content).json()["document_id"]

    response = client.get(f"{LIST_URL}/{document_id}/download")

    assert response.status_code == 200
    assert response.content == content  # byte-for-byte aynı dosya
    assert response.headers["content-type"] == media_type
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert document_id not in disposition  # storage adı (UUID) sızmaz


def test_list_and_detail_return_image_records_without_file_reference(client, fake_classify, fake_image_ocr):
    fake_image_ocr()
    fake_classify(classification())
    document_id = upload(client, "foto.png", make_image(image_format="png")).json()["document_id"]

    listed = client.get(LIST_URL).json()
    detail = client.get(f"{LIST_URL}/{document_id}").json()

    assert [item["file_type"] for item in listed] == ["png"]
    assert all("file_reference" not in item and "extracted_text" not in item for item in listed)
    assert detail["file_type"] == "png"
    assert detail["extracted_text"] == IMAGE_TEXT
    assert "file_reference" not in detail


# --- Legacy DOC (Word 97-2003) desteği ---

DOC_FIXTURE = pathlib.Path(__file__).parent / "fixtures" / "ornek_dilekce.doc"
DOC_KISA_FIXTURE = pathlib.Path(__file__).parent / "fixtures" / "ornek_kisa.doc"
DOC_TEXT_PARCASI = "Sokağımızdaki çöp konteynerlerinin boşaltılmaması"


def doc_content() -> bytes:
    return DOC_FIXTURE.read_bytes()


def test_doc_upload_is_classified(client, session_factory, storage_dir, fake_classify, caplog):
    calls = fake_classify(classification())

    with caplog.at_level(logging.DEBUG):
        response = upload(client, "dilekçe raporu.doc", doc_content())

    assert response.status_code == 200
    body = response.json()
    assert set(body) == RESPONSE_FIELDS  # file_reference ve extracted_text yanıtta yok
    assert body["file_type"] == "doc"
    assert body["file_name"] == "dilekçe raporu.doc"
    assert body["status"] == "classified"
    assert (body["summary"], body["sender_name"], body["sender_institution"]) == (
        SUMMARY, SENDER_NAME, SENDER_INSTITUTION,
    )
    assert len(calls) == 1  # belge başına tek sınıflandırma işlemi
    assert DOC_TEXT_PARCASI in calls[0]
    [document] = all_documents(session_factory)
    assert DOC_TEXT_PARCASI in document.extracted_text
    assert document.file_reference == f"{document.id}.doc"
    assert stored_files(storage_dir) == [f"{document.id}.doc"]


def test_doc_without_enough_text_is_failed_with_422(client, session_factory):
    response = upload(client, "kisa.doc", DOC_KISA_FIXTURE.read_bytes())

    assert response.status_code == 422
    body = response.json()
    assert body["status"] == "failed"
    assert body["message"] == TEXT_FAILED_MESSAGE
    assert body["document_type"] is None and body["institution_id"] is None
    assert all_documents(session_factory)[0].status == "failed"


@pytest.mark.parametrize(
    "file_name, content",
    [
        ("sahte.doc", b"duz metin dosyasi"),
        ("sahte.doc", make_pdf(PDF_TEXT)),
        ("sahte.doc", make_docx(DOCX_TEXT)),
    ],
    ids=["plain-text", "pdf-bytes", "docx-bytes"],
)
def test_fake_doc_returns_415(client, session_factory, storage_dir, file_name, content):
    response = upload(client, file_name, content)

    assert response.status_code == 415
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_oversized_doc_returns_413(client, session_factory, storage_dir, monkeypatch):
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", 1000)

    response = upload(client, "buyuk.doc", doc_content())

    assert response.status_code == 413
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_doc_download_returns_original_bytes_with_msword_media_type(client, fake_classify):
    fake_classify(classification())
    content = doc_content()
    document_id = upload(client, "dilekçe raporu.doc", content).json()["document_id"]

    response = client.get(f"{LIST_URL}/{document_id}/download")

    assert response.status_code == 200
    assert response.content == content  # byte-for-byte aynı dosya
    assert response.headers["content-type"] == "application/msword"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert "dilek" in disposition
    assert document_id not in disposition  # UUID storage adı sızmaz


def test_list_and_detail_return_doc_records_without_file_reference(client, fake_classify):
    fake_classify(classification())
    document_id = upload(client, "dilekce.doc", doc_content()).json()["document_id"]

    listed = client.get(LIST_URL).json()
    detail = client.get(f"{LIST_URL}/{document_id}").json()

    assert [item["file_type"] for item in listed] == ["doc"]
    assert all("file_reference" not in item and "extracted_text" not in item for item in listed)
    assert detail["file_type"] == "doc"
    assert DOC_TEXT_PARCASI in detail["extracted_text"]
    assert "file_reference" not in detail


# --- V1.4: hazırla → önizle → sınıflandır (D-045, D-046) ---

PREPARE_URL = "/api/documents/prepare"
DETAIL_FIELDS = RESPONSE_FIELDS | {"created_at", "extracted_text"}
ALREADY_PROCESSED_MESSAGE = "Belge zaten işlenmiş."
SOURCE_FILE_MISSING_MESSAGE = "Belgenin orijinal dosyasına ulaşılamadı; belgeyi kaldırıp yeniden yükleyin."


def prepare(client: TestClient, file_name: str, content: bytes):
    return client.post(PREPARE_URL, files={"file": (file_name, content, "application/octet-stream")})


def stored_document(session_factory, document_id) -> Document | None:
    with session_factory() as session:
        return session.get(Document, uuid.UUID(str(document_id)))


@pytest.fixture
def failing_commit(session_factory):
    """Kurulunca veritabanı commit'lerini düşürür; rollback çağrılarını kaydeden listeyi döndürür."""
    rollbacks = []

    def failing_get_db():
        with session_factory() as session:
            original_rollback = session.rollback

            def fail_commit():
                raise OperationalError("COMMIT", None, Exception("veritabanı erişilemiyor"))

            def tracking_rollback():
                rollbacks.append(True)
                original_rollback()

            session.commit = fail_commit
            session.rollback = tracking_rollback
            yield session

    def install():
        app.dependency_overrides[get_db] = failing_get_db
        return rollbacks

    return install


@pytest.mark.parametrize(
    "file_name, content, file_type",
    [
        ("dilekce.pdf", make_pdf(PDF_TEXT), "pdf"),
        ("dilekce.docx", make_docx(DOCX_TEXT), "docx"),
        ("dilekçe raporu.doc", DOC_FIXTURE.read_bytes(), "doc"),
        ("foto.png", make_image(image_format="png"), "png"),
    ],
    ids=["pdf", "docx", "doc", "png"],
)
def test_prepare_stores_file_and_text_without_calling_gemini(
    client, session_factory, storage_dir, fake_image_ocr, caplog, file_name, content, file_type
):
    # Gemini çağrılırsa autouse gemini_must_not_be_called_unexpectedly fixture'ı testi düşürür.
    fake_image_ocr()

    with caplog.at_level(logging.DEBUG):
        response = prepare(client, file_name, content)

    assert response.status_code == 200
    body = response.json()
    assert set(body) == DETAIL_FIELDS  # file_reference yanıtta yok
    assert (body["status"], body["file_name"], body["file_type"]) == ("prepared", file_name, file_type)
    assert (
        body["document_type"], body["document_type_name"], body["institution_id"], body["institution_name"],
        body["needs_review"], body["review_reason"], body["summary"], body["sender_name"], body["sender_institution"],
    ) == (None, None, None, None, False, None, None, None, None)
    [document] = all_documents(session_factory)
    assert str(document.id) == body["document_id"]
    assert document.status == "prepared"
    assert document.extracted_text == body["extracted_text"]
    assert len(document.extracted_text) >= file_service.MIN_TEXT_LENGTH
    assert (document.document_type, document.institution_id, document.review_reason, document.summary) == (None,) * 4
    assert document.file_reference == f"{document.id}.{file_type}"
    assert (storage_dir / document.file_reference).read_bytes() == content
    assert document.file_reference not in response.text
    assert SECRET_MARKER not in caplog.text


def test_prepare_rejects_oversized_file_without_record_or_file(client, session_factory, storage_dir, monkeypatch):
    content = make_pdf(PDF_TEXT)
    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", len(content) - 1)

    response = prepare(client, "buyuk.pdf", content)

    assert response.status_code == 413
    assert response.json() == {"detail": "Dosya boyutu 50 MB sınırını aşıyor."}
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


@pytest.mark.parametrize(
    "file_name, content",
    [("dilekce.txt", b"duz metin dosyasi"), ("dilekce.pdf", b"PDF olmayan icerik"), ("sahte.png", b"duz metin")],
    ids=["txt", "fake-pdf", "fake-png"],
)
def test_prepare_rejects_unsupported_file_without_record_or_file(client, session_factory, storage_dir, file_name, content):
    response = prepare(client, file_name, content)

    assert response.status_code == 415
    assert response.json() == {"detail": "Yalnızca PDF, DOC, DOCX, JPG, JPEG veya PNG dosyaları kabul edilir."}
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_prepare_without_file_returns_validation_error(client, session_factory, storage_dir):
    response = client.post(PREPARE_URL)

    assert response.status_code == 422
    assert "document_id" not in response.json()
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


@pytest.mark.parametrize(
    "file_name, content, expected_text",
    [
        ("kisa.docx", make_docx("Kısa"), "Kısa"),
        ("metinsiz.pdf", make_pdf(""), None),
        ("bozuk.pdf", b"%PDF-1.7\nbozuk icerik", None),
    ],
    ids=["short-text", "no-text", "corrupt"],
)
def test_prepare_text_failure_returns_422_with_failed_record(client, session_factory, storage_dir, file_name, content, expected_text):
    response = prepare(client, file_name, content)

    assert response.status_code == 422
    body = response.json()
    assert set(body) == RESPONSE_FIELDS | {"message"}
    assert (body["status"], body["message"]) == ("failed", TEXT_FAILED_MESSAGE)
    [document] = all_documents(session_factory)
    assert str(document.id) == body["document_id"]
    assert (document.status, document.document_type, document.needs_review) == ("failed", None, False)
    assert document.extracted_text == expected_text
    assert (storage_dir / document.file_reference).read_bytes() == content  # D-004: failed kaydının dosyası kalır


def test_prepare_storage_write_failure_returns_500_without_partial_file_or_record(
    client, session_factory, storage_dir, failing_storage_write
):
    response = prepare(client, "dilekce.pdf", make_pdf(PDF_TEXT))

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_prepare_commit_failure_rolls_back_and_removes_file(client, session_factory, storage_dir, failing_commit, caplog):
    rollbacks = failing_commit()

    with caplog.at_level(logging.DEBUG):
        response = prepare(client, "dilekce.pdf", make_pdf(PDF_TEXT))

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
    assert rollbacks == [True]
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []
    assert SECRET_MARKER not in caplog.text


def classify_url(document_id) -> str:
    return f"{LIST_URL}/{document_id}/classify"


def insert_with_file(session_factory, storage_dir, **overrides) -> Document:
    """Kayıt + storage dosyası (yerinde duran orijinal belge)."""
    document = insert_document(session_factory, **overrides)
    storage_dir.mkdir(parents=True, exist_ok=True)
    (storage_dir / document.file_reference).write_bytes(make_pdf(PDF_TEXT))
    return document


def insert_prepared(session_factory, storage_dir, **overrides) -> Document:
    """Hazırlanmış kayıt (D-046): metin var, sınıflandırma alanları boş; storage dosyası yerinde."""
    fields = {"status": "prepared", "document_type": None, "institution_id": None, "extracted_text": PDF_TEXT}
    return insert_with_file(session_factory, storage_dir, **(fields | overrides))


@pytest.fixture
def forbid_text_extraction(monkeypatch):
    """Kurulunca file_service.extract_text'i çağrıları kaydeden ve testi düşüren bir spy ile değiştirir."""
    def install():
        calls = []

        def extract_text(content, file_type):
            calls.append(file_type)
            raise AssertionError("metin çıkarımı/OCR sınıflandırma adımında tekrar çalıştı")

        monkeypatch.setattr(file_service, "extract_text", extract_text)
        return calls

    return install


@pytest.mark.parametrize("needs_review", [False, True], ids=["classified", "needs_review"])
def test_classify_prepared_document_uses_stored_text_without_reextracting(
    client, session_factory, storage_dir, fake_classify, forbid_text_extraction, caplog, needs_review
):
    content = make_pdf(PDF_TEXT)
    document_id = prepare(client, "dilekce.pdf", content).json()["document_id"]
    extraction_calls = forbid_text_extraction()  # storage dosyası yerinde; yalnız yeniden çıkarım yasak
    classify_calls = fake_classify(classification(needs_review=needs_review))

    with caplog.at_level(logging.DEBUG):
        response = client.post(classify_url(document_id))

    assert response.status_code == 200
    body = response.json()
    assert set(body) == RESPONSE_FIELDS
    expected_status = "needs_review" if needs_review else "classified"
    assert (body["document_id"], body["status"], body["summary"]) == (document_id, expected_status, SUMMARY)
    assert classify_calls == [PDF_TEXT]  # DB'deki metinle tek Gemini çağrısı
    assert extraction_calls == []  # OCR/extraction tekrar çalışmadı
    document = stored_document(session_factory, document_id)
    assert (document.status, document.summary, document.extracted_text) == (expected_status, SUMMARY, PDF_TEXT)
    assert (storage_dir / document.file_reference).read_bytes() == content
    assert SECRET_MARKER not in caplog.text


def test_classify_prepared_document_gemini_error_returns_502_and_failed_record(
    client, session_factory, storage_dir, fake_classify, caplog
):
    document_id = prepare(client, "dilekce.docx", make_docx(DOCX_TEXT)).json()["document_id"]
    fake_classify(classification_error())

    with caplog.at_level(logging.DEBUG):
        response = client.post(classify_url(document_id))

    assert response.status_code == 502
    body = response.json()
    assert set(body) == RESPONSE_FIELDS | {"message"}
    assert (body["status"], body["message"]) == ("failed", CLASSIFICATION_FAILED_MESSAGE)
    assert "ham Gemini" not in response.text and "API_KEY_INVALID" not in response.text
    document = stored_document(session_factory, document_id)
    assert (document.status, document.document_type, document.institution_id, document.needs_review, document.summary) == (
        "failed", None, None, False, None,
    )
    assert document.extracted_text == DOCX_TEXT  # D-018: metin korunur
    assert (storage_dir / document.file_reference).exists()
    assert SECRET_MARKER not in caplog.text


def test_classify_unknown_document_returns_404(client, fake_classify):
    calls = fake_classify(classification())

    response = client.post(classify_url(uuid.uuid4()))

    assert response.status_code == 404
    assert response.json() == {"detail": "Belge bulunamadı."}
    assert calls == []


def test_classify_with_invalid_id_returns_validation_error(client, fake_classify):
    calls = fake_classify(classification())

    assert client.post(classify_url("belge-degil")).status_code == 422
    assert calls == []


@pytest.mark.parametrize("status", ["classified", "needs_review", "failed"])
def test_classify_rejects_non_prepared_document_with_409(client, session_factory, storage_dir, fake_classify, status):
    document = insert_with_file(session_factory, storage_dir, status=status)
    calls = fake_classify(classification())

    response = client.post(classify_url(document.id))

    assert response.status_code == 409
    assert response.json() == {"detail": ALREADY_PROCESSED_MESSAGE}
    assert calls == []
    assert stored_document(session_factory, document.id).status == status


def test_classify_commit_failure_keeps_prepared_record_and_file(client, session_factory, storage_dir, fake_classify, failing_commit):
    document = insert_prepared(session_factory, storage_dir)
    fake_classify(classification())
    rollbacks = failing_commit()

    response = client.post(classify_url(document.id))

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
    assert rollbacks == [True]
    stored = stored_document(session_factory, document.id)
    assert (stored.status, stored.document_type, stored.summary) == ("prepared", None, None)
    assert (storage_dir / document.file_reference).exists()  # dosya silinmez; kayıt tekrar denenebilir


def test_classify_without_source_file_returns_409_and_keeps_prepared_record(
    client, session_factory, storage_dir, fake_classify, forbid_text_extraction, caplog
):
    document_id = prepare(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()["document_id"]
    document = stored_document(session_factory, document_id)
    (storage_dir / document.file_reference).unlink()  # discard yarım kaldı veya dosya kayboldu
    classify_calls = fake_classify(classification())
    extraction_calls = forbid_text_extraction()

    with caplog.at_level(logging.DEBUG):
        response = client.post(classify_url(document_id))

    assert response.status_code == 409
    assert response.json() == {"detail": SOURCE_FILE_MISSING_MESSAGE}
    assert classify_calls == []  # Gemini 0 çağrı
    assert extraction_calls == []  # extraction 0 çağrı
    after = stored_document(session_factory, document_id)
    assert (after.status, after.extracted_text, after.document_type, after.summary) == ("prepared", PDF_TEXT, None, None)
    # Storage yolu ve belge metni ne yanıtta ne logda.
    assert document.file_reference not in response.text and str(storage_dir) not in response.text
    assert document.file_reference not in caplog.text and str(storage_dir) not in caplog.text
    assert SECRET_MARKER not in caplog.text


def test_classify_rejects_file_reference_outside_storage(
    client, session_factory, storage_dir, tmp_path, fake_classify, forbid_text_extraction, caplog
):
    outside = tmp_path / "disarida.pdf"
    outside.write_bytes(make_pdf(PDF_TEXT))
    storage_dir.mkdir(parents=True, exist_ok=True)
    document = insert_document(
        session_factory, status="prepared", document_type=None, institution_id=None, file_reference="../disarida.pdf",
    )
    classify_calls = fake_classify(classification())
    extraction_calls = forbid_text_extraction()

    with caplog.at_level(logging.DEBUG):
        response = client.post(classify_url(document.id))

    assert response.status_code == 409
    assert response.json() == {"detail": SOURCE_FILE_MISSING_MESSAGE}
    assert classify_calls == [] and extraction_calls == []
    assert "disarida" not in response.text and "disarida" not in caplog.text
    assert stored_document(session_factory, document.id).status == "prepared"
    assert outside.exists()


def discard_url(document_id) -> str:
    return f"{LIST_URL}/{document_id}/prepared"


def test_discard_prepared_document_deletes_record_and_file(client, session_factory, storage_dir, caplog):
    document_id = prepare(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()["document_id"]
    assert len(stored_files(storage_dir)) == 1

    with caplog.at_level(logging.DEBUG):
        response = client.delete(discard_url(document_id))

    assert response.status_code == 204
    assert response.content == b""
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []
    assert SECRET_MARKER not in caplog.text


@pytest.mark.parametrize("status", ["classified", "needs_review", "failed"])
def test_discard_rejects_permanent_records_with_409(client, session_factory, storage_dir, fake_classify, status):
    document = insert_with_file(session_factory, storage_dir, status=status)
    calls = fake_classify(classification())

    response = client.delete(discard_url(document.id))

    assert response.status_code == 409
    assert response.json() == {"detail": ALREADY_PROCESSED_MESSAGE}
    assert stored_document(session_factory, document.id).status == status
    assert (storage_dir / document.file_reference).exists()
    assert calls == []


def test_discard_unknown_document_returns_404(client):
    response = client.delete(discard_url(uuid.uuid4()))

    assert response.status_code == 404
    assert response.json() == {"detail": "Belge bulunamadı."}


def test_discard_with_invalid_id_returns_validation_error(client):
    assert client.delete(discard_url("belge-degil")).status_code == 422


def test_discard_works_when_storage_file_is_already_missing(client, session_factory, storage_dir):
    document = insert_prepared(session_factory, storage_dir)
    (storage_dir / document.file_reference).unlink()

    response = client.delete(discard_url(document.id))

    assert response.status_code == 204
    assert all_documents(session_factory) == []


def test_second_discard_of_same_document_returns_404(client, session_factory, storage_dir):
    document = insert_prepared(session_factory, storage_dir)

    assert client.delete(discard_url(document.id)).status_code == 204
    assert client.delete(discard_url(document.id)).status_code == 404


def test_discard_file_delete_error_keeps_prepared_record_for_retry(client, session_factory, storage_dir, monkeypatch):
    document = insert_prepared(session_factory, storage_dir)
    original_unlink = pathlib.Path.unlink

    def locked_unlink(self, *args, **kwargs):
        raise PermissionError("dosya kilitli")

    monkeypatch.setattr(pathlib.Path, "unlink", locked_unlink)
    response = client.delete(discard_url(document.id))

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
    assert stored_document(session_factory, document.id).status == "prepared"
    assert (storage_dir / document.file_reference).exists()

    monkeypatch.setattr(pathlib.Path, "unlink", original_unlink)
    assert client.delete(discard_url(document.id)).status_code == 204
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_discard_commit_failure_keeps_prepared_record_for_retry(client, session_factory, storage_dir, failing_commit):
    document = insert_prepared(session_factory, storage_dir)
    rollbacks = failing_commit()

    response = client.delete(discard_url(document.id))

    assert response.status_code == 500
    assert rollbacks == [True]
    assert stored_document(session_factory, document.id).status == "prepared"  # TTL ve tekrar Kaldır bulabilir

    def working_get_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = working_get_db
    assert client.delete(discard_url(document.id)).status_code == 204  # dosya zaten silinmiş olsa da güvenli
    assert all_documents(session_factory) == []
    assert stored_files(storage_dir) == []


def test_discard_never_deletes_files_outside_storage(client, session_factory, storage_dir, tmp_path):
    outside = tmp_path / "disarida.pdf"
    outside.write_bytes(make_pdf(PDF_TEXT))
    storage_dir.mkdir(parents=True, exist_ok=True)
    document = insert_document(
        session_factory, status="prepared", document_type=None, institution_id=None, file_reference="../disarida.pdf",
    )

    response = client.delete(discard_url(document.id))

    assert response.status_code == 204
    assert all_documents(session_factory) == []
    assert outside.exists()


def test_discarded_document_is_gone_from_detail_download_and_classify(client, storage_dir, fake_classify):
    document_id = prepare(client, "dilekce.pdf", make_pdf(PDF_TEXT)).json()["document_id"]
    calls = fake_classify(classification())

    assert client.delete(discard_url(document_id)).status_code == 204

    assert client.get(f"{LIST_URL}/{document_id}").status_code == 404
    assert client.get(f"{LIST_URL}/{document_id}/download").status_code == 404
    assert client.post(classify_url(document_id)).status_code == 404
    assert calls == []


def test_openapi_documents_v14_endpoints():
    paths = app.openapi()["paths"]

    prepare_responses = paths[PREPARE_URL]["post"]["responses"]
    assert {"200", "413", "415", "422"} <= set(prepare_responses) and "502" not in prepare_responses
    classify_responses = paths[f"{LIST_URL}/{{document_id}}/classify"]["post"]["responses"]
    assert {"200", "404", "409", "502"} <= set(classify_responses)
    discard_responses = paths[f"{LIST_URL}/{{document_id}}/prepared"]["delete"]["responses"]
    assert {"204", "404", "409"} <= set(discard_responses)


def test_list_hides_prepared_documents_but_detail_and_download_work(client, session_factory):
    content = make_pdf(PDF_TEXT)
    prepared_id = prepare(client, "hazir.pdf", content).json()["document_id"]
    insert_document(session_factory, file_name="kalici.pdf")

    assert [item["file_name"] for item in client.get(LIST_URL).json()] == ["kalici.pdf"]
    detail = client.get(f"{LIST_URL}/{prepared_id}")
    assert detail.status_code == 200
    assert (detail.json()["status"], detail.json()["extracted_text"]) == ("prepared", PDF_TEXT)
    download = client.get(f"{LIST_URL}/{prepared_id}/download")
    assert download.status_code == 200
    assert download.content == content


# Yedek temizlik (TTL): 24 saatten eski sahipsiz prepared kayıtlar sonraki prepare çağrısında silinir (D-046).

def hours_ago(hours: int) -> datetime:
    return datetime.now(timezone.utc) - timedelta(hours=hours)


def test_prepare_cleans_up_expired_orphan_prepared_documents(client, session_factory, storage_dir):
    expired = insert_prepared(session_factory, storage_dir, created_at=hours_ago(25))

    new_id = prepare(client, "yeni.pdf", make_pdf(PDF_TEXT)).json()["document_id"]

    assert [str(document.id) for document in all_documents(session_factory)] == [new_id]
    assert not (storage_dir / expired.file_reference).exists()


def test_cleanup_keeps_fresh_prepared_and_all_permanent_records(client, session_factory, storage_dir):
    fresh = insert_prepared(session_factory, storage_dir, created_at=hours_ago(1))
    permanent = [
        insert_with_file(session_factory, storage_dir, status=status, created_at=hours_ago(25))
        for status in ("classified", "needs_review", "failed")
    ]

    new_id = prepare(client, "yeni.pdf", make_pdf(PDF_TEXT)).json()["document_id"]

    kept = [fresh, *permanent]
    assert {str(document.id) for document in all_documents(session_factory)} == {new_id, *(str(d.id) for d in kept)}
    assert all((storage_dir / document.file_reference).exists() for document in kept)


def test_cleanup_skips_record_whose_file_cannot_be_deleted(client, session_factory, storage_dir, monkeypatch, caplog):
    stuck = insert_prepared(session_factory, storage_dir, created_at=hours_ago(25))
    other = insert_prepared(session_factory, storage_dir, created_at=hours_ago(25))
    stuck_path = (storage_dir / stuck.file_reference).resolve()
    original_unlink = pathlib.Path.unlink

    def unlink(self, *args, **kwargs):
        if self == stuck_path:
            raise PermissionError("dosya kilitli")
        return original_unlink(self, *args, **kwargs)

    monkeypatch.setattr(pathlib.Path, "unlink", unlink)
    with caplog.at_level(logging.DEBUG):
        response = prepare(client, "yeni.pdf", make_pdf(PDF_TEXT))

    assert response.status_code == 200
    remaining = {document.id for document in all_documents(session_factory)}
    assert stuck.id in remaining and other.id not in remaining  # dosyası silinemeyen kayıt sonraki temizliğe kalır
    assert stuck_path.exists()
    assert not (storage_dir / other.file_reference).exists()
    assert stuck.file_reference not in caplog.text


def test_cleanup_failure_does_not_break_prepare(client, session_factory, storage_dir, monkeypatch):
    from app.api import documents as documents_api

    expired = insert_prepared(session_factory, storage_dir, created_at=hours_ago(25))

    def broken_discard(document, db):
        raise RuntimeError("beklenmeyen temizlik hatası")

    monkeypatch.setattr(documents_api, "_discard_prepared", broken_discard)
    response = prepare(client, "yeni.pdf", make_pdf(PDF_TEXT))

    assert response.status_code == 200
    assert response.json()["status"] == "prepared"
    assert stored_document(session_factory, expired.id).status == "prepared"


def test_expired_prepared_document_is_only_cleaned_by_prepare(client, session_factory, storage_dir, fake_classify):
    """Zamanlayıcı yok: liste, detay, indirme, classify ve discard temizlik tetiklemez."""
    expired = insert_prepared(session_factory, storage_dir, created_at=hours_ago(25))
    to_classify = insert_prepared(session_factory, storage_dir)
    to_discard = insert_prepared(session_factory, storage_dir)
    fake_classify(classification())

    client.get(LIST_URL)
    client.get(f"{LIST_URL}/{expired.id}")
    client.get(f"{LIST_URL}/{expired.id}/download")
    assert client.post(classify_url(to_classify.id)).status_code == 200
    assert client.delete(discard_url(to_discard.id)).status_code == 204

    assert stored_document(session_factory, expired.id).status == "prepared"
    assert (storage_dir / expired.file_reference).exists()

    prepare(client, "yeni.pdf", make_pdf(PDF_TEXT))
    assert stored_document(session_factory, expired.id) is None

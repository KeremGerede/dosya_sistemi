import io
import logging
import uuid
import zipfile

import docx
import pymupdf
import pytest

from app.services import file_service
from app.services.file_service import (
    MAX_FILE_SIZE,
    MIN_TEXT_LENGTH,
    FileTooLargeError,
    TextExtractionError,
    UnsupportedFileTypeError,
    check_file_size,
    check_text_length,
    detect_file_type,
    extract_text,
    normalize_text,
    save_file,
)

DOCX_CONTENT_TYPES = (
    b'<?xml version="1.0" encoding="UTF-8"?>'
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Override PartName="/word/document.xml" '
    b'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    b"</Types>"
)


def make_pdf(*pages: str) -> bytes:
    with pymupdf.open() as pdf:
        for text in pages:
            page = pdf.new_page()
            if text:
                page.insert_text((72, 72), text)
        return pdf.tobytes()


def docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx(*paragraphs: str) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return docx_bytes(document)


def make_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.fixture
def storage_dir(tmp_path, monkeypatch):
    """Testler gerçek backend/storage yerine geçici bir klasöre yazar."""
    path = tmp_path / "storage"
    monkeypatch.setattr(file_service, "STORAGE_DIR", path)
    return path


def test_valid_pdf_is_accepted_and_text_extracted():
    data = make_pdf("Birinci sayfa metni", "Ikinci sayfa metni")

    assert detect_file_type("dilekce.pdf", data) == "pdf"
    assert detect_file_type("DILEKCE.PDF", data) == "pdf"
    text = extract_text(data, "pdf")
    assert text == "Birinci sayfa metni Ikinci sayfa metni"
    check_text_length(text)


def test_valid_docx_is_accepted_and_paragraph_text_extracted():
    data = make_docx("Çöp konteyneri   taşmış durumda.", "", "Gereğinin yapılmasını\trica ederim.")

    assert detect_file_type("basvuru.docx", data) == "docx"
    assert extract_text(data, "docx") == "Çöp konteyneri taşmış durumda. Gereğinin yapılmasını rica ederim."


def test_docx_table_cell_text_is_extracted_once_in_document_order():
    document = docx.Document()
    document.add_paragraph("Başvuru bilgileri:")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Birleşik başlık"
    table.cell(1, 0).text = "Ad Soyad"
    table.cell(1, 1).text = "Ayşe Yılmaz"
    document.add_paragraph("Son paragraf.")

    text = extract_text(docx_bytes(document), "docx")

    assert text == "Başvuru bilgileri: Birleşik başlık Ad Soyad Ayşe Yılmaz Son paragraf."
    assert text.count("Birleşik başlık") == 1


def test_text_shorter_than_10_characters_is_rejected():
    check_text_length("1234567890")
    with pytest.raises(TextExtractionError):
        check_text_length("123456789")
    with pytest.raises(TextExtractionError):
        check_text_length(normalize_text("  a   b  c  "))

    # Görsel/taranmış PDF gibi metin vermeyen sayfa: OCR yok, yetersiz metin.
    blank_pdf_text = extract_text(make_pdf(""), "pdf")
    assert blank_pdf_text == ""
    with pytest.raises(TextExtractionError):
        check_text_length(blank_pdf_text)

    with pytest.raises(TextExtractionError):
        check_text_length(extract_text(make_docx("Kısa"), "docx"))


@pytest.mark.parametrize(
    "file_name, content_factory",
    [
        ("dilekce.doc", lambda: make_docx("Word 97 uzantılı dosya")),
        ("dilekce.txt", lambda: make_pdf("PDF icerik")),
        ("dilekce", lambda: make_pdf("PDF icerik")),
        ("dilekce.pdf.exe", lambda: make_pdf("PDF icerik")),
        ("dilekce.pdf", lambda: make_docx("DOCX içerik .pdf uzantılı")),
        ("dilekce.docx", lambda: make_pdf("PDF icerik")),
    ],
)
def test_wrong_extension_or_mismatched_content_is_rejected(file_name, content_factory):
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(file_name, content_factory())


@pytest.mark.parametrize("content", [b"Merhaba, ben PDF degilim", b"", b" %PDF-1.7 basta bosluk var"])
def test_pdf_extension_with_non_pdf_content_is_rejected(content):
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type("dilekce.pdf", content)


@pytest.mark.parametrize(
    "content",
    [
        b"PK ama zip degil",
        make_zip({"readme.txt": b"generic zip"}),
        make_zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<w:document/>"}),
        make_zip({"word/document.xml": b"<w:document/>"}),
        make_zip({"[Content_Types].xml": DOCX_CONTENT_TYPES}),
    ],
    ids=["not-zip", "generic-zip", "no-docx-content-type", "no-content-types", "no-document-xml"],
)
def test_docx_extension_with_non_docx_content_is_rejected(content):
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type("dilekce.docx", content)


def test_corrupt_or_encrypted_documents_raise_text_extraction_error():
    corrupt_pdf = b"%PDF-1.7\nbu dosya bozuk"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Gizli dilekce metni")
        encrypted_pdf = pdf.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="kullanici", owner_pw="sahip")
    corrupt_docx = make_zip({"[Content_Types].xml": DOCX_CONTENT_TYPES, "word/document.xml": b"<bozuk"})

    for content, file_type in [(corrupt_pdf, "pdf"), (encrypted_pdf, "pdf"), (corrupt_docx, "docx")]:
        assert detect_file_type(f"belge.{file_type}", content) == file_type
        with pytest.raises(TextExtractionError):
            check_text_length(extract_text(content, file_type))


def test_storage_file_name_is_built_from_document_id(storage_dir):
    document_id = uuid.uuid4()
    data = make_pdf("Dosya icerigi")

    file_reference = save_file(data, document_id, "pdf")

    assert file_reference == f"{document_id}.pdf"
    assert (storage_dir / file_reference).read_bytes() == data
    assert [path.name for path in storage_dir.iterdir()] == [file_reference]


def test_user_file_name_is_not_used_as_storage_path(storage_dir, tmp_path):
    user_file_name = "../../gizli/dilekce.pdf"
    data = make_pdf("Dosya icerigi")
    document_id = uuid.uuid4()

    file_reference = save_file(data, document_id, detect_file_type(user_file_name, data))

    assert file_reference == f"{document_id}.pdf"
    assert "dilekce" not in file_reference and ".." not in file_reference and "/" not in file_reference
    assert list(tmp_path.rglob("*.pdf")) == [storage_dir / file_reference]
    with pytest.raises(ValueError):
        save_file(data, "../../gizli", "pdf")
    with pytest.raises(ValueError):
        save_file(data, uuid.uuid4(), "../pdf")


def test_failed_write_removes_partial_file_and_reraises(storage_dir, failing_storage_write):
    with pytest.raises(OSError) as exc_info:
        save_file(make_pdf("Dosya icerigi"), uuid.uuid4(), "pdf")

    assert exc_info.value is failing_storage_write  # özgün hata bastırılmaz
    assert list(storage_dir.iterdir()) == []  # yarım yazılmış dosya kalmaz


def test_existing_storage_file_is_never_overwritten_or_deleted(storage_dir):
    document_id = uuid.uuid4()
    original = make_pdf("Ilk dosya")
    file_reference = save_file(original, document_id, "pdf")

    with pytest.raises(FileExistsError):
        save_file(make_pdf("Ikinci dosya"), document_id, "pdf")

    assert (storage_dir / file_reference).read_bytes() == original


def test_full_text_is_returned_without_truncation():
    long_text = "a" * 60_000

    assert extract_text(make_docx(long_text), "docx") == long_text


def test_normalize_text_collapses_whitespace_and_trims():
    assert normalize_text("  Merhaba\n\n\tdünya   nasılsın?  ") == "Merhaba dünya nasılsın?"
    assert normalize_text(" \n\t ") == ""


def test_file_size_limit_is_50_mb():
    assert MAX_FILE_SIZE == 50 * 1024 * 1024
    check_file_size(0)
    check_file_size(MAX_FILE_SIZE)
    with pytest.raises(FileTooLargeError):
        check_file_size(MAX_FILE_SIZE + 1)


# --- Taranmış PDF için OCR fallback (V1.1) ---


def test_pdf_with_enough_text_does_not_use_ocr(monkeypatch):
    def fail_if_called(page):
        raise AssertionError("Yeterli metni olan sayfada OCR çağrılmamalı.")

    monkeypatch.setattr(file_service, "_ocr_page_text", fail_if_called)
    assert extract_text(make_pdf("Bu dilekçe yeterli uzunlukta metin içerir."), "pdf") == (
        "Bu dilekçe yeterli uzunlukta metin içerir."
    )


def test_pdf_without_enough_text_uses_ocr_fallback(monkeypatch):
    monkeypatch.setattr(file_service, "_ocr_page_text", lambda page: "OCR ile okunan şikayet dilekçesi.")
    assert extract_text(make_pdf(""), "pdf") == "OCR ile okunan şikayet dilekçesi."


def test_ocr_output_long_enough_passes_length_check(monkeypatch):
    monkeypatch.setattr(file_service, "_ocr_page_text", lambda page: "OCR ile okunan yeterli metin.")
    check_text_length(extract_text(make_pdf(""), "pdf"))  # hata yükselmemeli


def test_ocr_output_still_too_short_is_rejected(monkeypatch):
    monkeypatch.setattr(file_service, "_ocr_page_text", lambda page: "kısa")
    text = extract_text(make_pdf(""), "pdf")
    assert text == "kısa"
    with pytest.raises(TextExtractionError):
        check_text_length(text)


def test_ocr_failure_falls_back_to_safe_text_extraction_failure(monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("tesseract dili yükleyemedi")

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", "/tessdata")
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", boom)
    text = extract_text(make_pdf(""), "pdf")  # OCR hatası dışarı sızmaz
    assert text == ""
    with pytest.raises(TextExtractionError):
        check_text_length(text)


def test_ocr_is_skipped_when_tessdata_prefix_is_not_configured(monkeypatch):
    def fail_if_called(*args, **kwargs):
        raise AssertionError("TESSDATA_PREFIX tanımlı değilken OCR denenmemeli.")

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", None)
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", fail_if_called)
    assert extract_text(make_pdf(""), "pdf") == ""


def test_ocr_is_called_with_turkish_and_400_dpi(monkeypatch):
    calls = []

    def fake_ocr(self, *args, **kwargs):
        calls.append(kwargs)
        return self.get_textpage()

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", "/tessdata")
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", fake_ocr)
    extract_text(make_pdf(""), "pdf")
    assert file_service.OCR_DPI == 400
    assert [(c["language"], c["dpi"], c["tessdata"]) for c in calls] == [("tur", 400, "/tessdata")]


def test_docx_never_uses_ocr(monkeypatch):
    def fail_if_called(page):
        raise AssertionError("DOCX için OCR çağrılmamalı.")

    monkeypatch.setattr(file_service, "_ocr_page_text", fail_if_called)
    with pytest.raises(TextExtractionError):
        check_text_length(extract_text(make_docx("kısa"), "docx"))


# --- Sayfa düzeyinde OCR kararı (V1.2 · P1 düzeltmesi) ---

OCR_PAGE_TEXT = "Bu sayfa OCR ile okunan taranmis dilekce metnidir."


def make_mixed_pdf(*pages: tuple[str, str]) -> bytes:
    """pages: ("text", içerik) gerçek text layer, ("image", içerik) metin katmanı olmayan taranmış sayfa."""
    with pymupdf.open() as pdf:
        for kind, content in pages:
            if kind == "text":
                page = pdf.new_page()
                if content:
                    page.insert_text((72, 72), content)
                continue
            with pymupdf.open() as source:  # önce çiz, sonra yalnızca görüntü olarak göm
                drawn = source.new_page()
                drawn.insert_text((72, 72), content)
                png = drawn.get_pixmap(dpi=72).tobytes("png")
            page = pdf.new_page()
            page.insert_image(page.rect, stream=png)
        return pdf.tobytes()


@pytest.fixture
def ocr_spy(monkeypatch):
    """_ocr_page_text yerine geçer; OCR'lanan sayfa numaralarını kaydeder."""
    pages = []

    def fake_ocr(page):
        pages.append(page.number)
        return OCR_PAGE_TEXT

    monkeypatch.setattr(file_service, "_ocr_page_text", fake_ocr)
    return pages


def test_all_text_pages_never_trigger_ocr(ocr_spy):
    content = make_mixed_pdf(("text", "Birinci sayfanin yeterli uzunlukta metni."),
                             ("text", "Ikinci sayfanin yeterli uzunlukta metni."))

    text = extract_text(content, "pdf")

    assert ocr_spy == []  # hiçbir sayfada OCR yok
    assert "Birinci sayfanin" in text and "Ikinci sayfanin" in text


def test_single_image_only_page_triggers_ocr(ocr_spy):
    """V1.1 davranışı korunur: metinsiz tek sayfalık PDF OCR'lanır."""
    text = extract_text(make_mixed_pdf(("image", "Taranmis dilekce metni")), "pdf")

    assert ocr_spy == [0]
    assert text == OCR_PAGE_TEXT


def test_hybrid_pdf_ocrs_only_the_image_page_and_keeps_order(ocr_spy):
    """P1: kapak sayfasının metni, taranmış ikinci sayfanın OCR'lanmasını engellememeli."""
    cover = "EVRAK KAYIT FORMU Bu belge resmi evrak kayit sistemine alinmistir."
    content = make_mixed_pdf(("text", cover), ("image", "Asil dilekce taranmis sayfada"))

    text = extract_text(content, "pdf")

    assert ocr_spy == [1]  # yalnızca 2. sayfa
    assert cover.split()[0] in text and OCR_PAGE_TEXT in text
    assert text.index("EVRAK") < text.index(OCR_PAGE_TEXT)  # belge sırası korunur


def test_multi_page_hybrid_ocrs_only_image_pages(ocr_spy):
    content = make_mixed_pdf(
        ("text", "Birinci sayfanin yeterli uzunlukta metni."),
        ("image", "Ikinci sayfa taranmis"),
        ("text", "Ucuncu sayfanin yeterli uzunlukta metni."),
        ("image", "Dorduncu sayfa taranmis"),
    )

    text = extract_text(content, "pdf")

    assert ocr_spy == [1, 3]  # yalnızca görüntü sayfaları, sırasıyla
    assert text.count(OCR_PAGE_TEXT) == 2
    assert "Birinci sayfanin" in text and "Ucuncu sayfanin" in text


@pytest.mark.parametrize(
    "page_text, ocr_expected",
    [("123456789", True), ("1234567890", False)],  # 9 karakter -> OCR, 10 karakter -> OCR yok
)
def test_page_level_threshold_is_min_text_length(ocr_spy, page_text, ocr_expected):
    assert len(page_text) == (MIN_TEXT_LENGTH - 1 if ocr_expected else MIN_TEXT_LENGTH)

    extract_text(make_mixed_pdf(("text", page_text)), "pdf")

    assert ocr_spy == ([0] if ocr_expected else [])


def test_hybrid_page_ocr_failure_falls_back_to_embedded_text(monkeypatch, caplog):
    """OCR hatası yeni hata sınıfı üretmez; sayfa gömülü metnine düşer, teknik detay sızmaz."""
    secret = "TARANMIS_SAYFA_ICERIGI"
    cover = "EVRAK KAYIT FORMU Bu belge resmi evrak kayit sistemine alinmistir."

    def boom(*args, **kwargs):
        raise RuntimeError(f"tesseract patladi: {secret}")

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", "/tessdata")
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", boom)

    with caplog.at_level(logging.WARNING):
        text = extract_text(make_mixed_pdf(("text", cover), ("image", secret)), "pdf")

    assert "EVRAK" in text  # kapak metni korunur, istek hata vermez
    assert secret not in text
    assert secret not in caplog.text  # ham hata mesajı ve belge metni loglanmaz
    assert "RuntimeError" in caplog.text  # güvenli bağlam: yalnızca hata türü


# --- Yapısal OCR fallback ve embedded/OCR birleştirme (V1.2 · P2) ---

SCAN_OCR_TEXT = (
    "KADIKOY BELEDIYE BASKANLIGINA Sokagimizdaki cop konteynerleri uzun suredir bosaltilmamakta, "
    "etrafa yayilan atiklar ve koku nedeniyle rahatsizlik yasanmaktadir. Geregini arz ederim."
)
# Bozuk text layer: hatalı ToUnicode CMap'ten gelen, kelime içermeyen ~30 karakter (senaryo 14 ailesi).
GARBAGE_LAYER = "¤¤ ¬¬ ¸¸ || ~~ ±± §§ ¶¶ ©© ®®"


def make_scan_pdf(image_text: str, layer: str = "", coverage: float = 1.0) -> bytes:
    """Taranmış sayfa: sayfanın `coverage` oranını kaplayan görüntü + isteğe bağlı gömülü metin katmanı."""
    with pymupdf.open() as source:
        drawn = source.new_page()
        drawn.insert_text((72, 72), image_text)
        png = drawn.get_pixmap(dpi=72).tobytes("png")
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        rect = page.rect
        page.insert_image(
            pymupdf.Rect(rect.x0, rect.y0, rect.x1, rect.y0 + rect.height * coverage),
            stream=png,
            keep_proportion=False,  # kapsama oranı testte birebir ölçülebilsin
        )
        for row, line in enumerate(layer.splitlines()):  # tek satır sayfa kenarında kesilir
            page.insert_text((72, 30 + row * 12), line, render_mode=3)  # tarayıcı katmanları görünmez yazılır
        return pdf.tobytes()


def layer_of_length(length: int) -> str:
    """Sayfaya sığan satırlara bölünmüş, normalize edildiğinde tam `length` karakter veren metin katmanı."""
    lines = -(-length // 50)
    chars = length - (lines - 1)  # normalize sırasında her satır sonu tek boşluğa döner
    size, extra = divmod(chars, lines)
    return "\n".join("a" * (size + (1 if row < extra else 0)) for row in range(lines))


@pytest.fixture
def fake_ocr(monkeypatch):
    """_ocr_page_text yerine geçer; OCR çağrılan sayfa numaralarını kaydeder, sabit metin döndürür."""
    def install(text=SCAN_OCR_TEXT):
        pages = []

        def fake(page):
            pages.append(page.number)
            return text

        monkeypatch.setattr(file_service, "_ocr_page_text", fake)
        return pages

    return install


def test_scanned_page_with_broken_text_layer_is_ocred(fake_ocr):
    """P2 / benchmark senaryo 14: 10+ karakterlik ama bozuk katman OCR'ı engellememeli."""
    pages = fake_ocr()
    content = make_scan_pdf("Taranmis dilekce metni", layer=GARBAGE_LAYER)

    assert len(normalize_text(GARBAGE_LAYER)) >= MIN_TEXT_LENGTH  # bugünkü eşiği geçiyor
    text = extract_text(content, "pdf")

    assert pages == [0]
    assert SCAN_OCR_TEXT in text


def test_broken_text_layer_does_not_pollute_ocr_text(fake_ocr):
    """M2: kelime içermeyen bozuk katman nihai metne taşınmaz."""
    fake_ocr()
    text = extract_text(make_scan_pdf("Taranmis dilekce metni", layer=GARBAGE_LAYER), "pdf")

    assert text == SCAN_OCR_TEXT
    assert "¤¤" not in text and "¶¶" not in text


def test_valuable_embedded_text_is_kept_when_ocr_does_not_contain_it(fake_ocr):
    """M1: OCR metni evrak bilgisini içermiyorsa gömülü metin korunur (bilgi kaybı yok)."""
    fake_ocr()
    text = extract_text(make_scan_pdf("Taranmis dilekce", layer="Evrak No: 2026/4417 Tarih: 18.09.2026"), "pdf")

    assert "2026/4417" in text  # gömülü evrak bilgisi
    assert "18.09.2026" in text
    assert SCAN_OCR_TEXT in text  # görüntüdeki dilekçe
    assert text.index("2026/4417") < text.index("KADIKOY")  # sayfa içi sıra korunur


def test_embedded_text_already_read_by_ocr_is_not_duplicated(fake_ocr):
    """M1 gerçek hayatta: OCR tüm sayfayı okuduğu için evrak bilgisini de görür; tekrar eklenmez."""
    fake_ocr("Evrak No: 2026/4417 Tarih: 18.09.2026 " + SCAN_OCR_TEXT)
    text = extract_text(make_scan_pdf("Taranmis dilekce", layer="Evrak No: 2026/4417 Tarih: 18.09.2026"), "pdf")

    assert text.count("2026/4417") == 1  # duplicate yok
    assert "18.09.2026" in text and SCAN_OCR_TEXT in text  # iki kaynağın bilgisi de var


def test_identical_embedded_and_ocr_content_is_written_once(fake_ocr):
    """M3: aynı içerik hem katmanda hem görüntüdeyse tek kez yazılır (Türkçe karakter farkına rağmen)."""
    fake_ocr("Sikayet Dilekcesi " + SCAN_OCR_TEXT)
    text = extract_text(make_scan_pdf("Sikayet Dilekcesi", layer="Sikayet Dilekçesi"), "pdf")

    assert text.lower().count("ikayet") == 1
    assert SCAN_OCR_TEXT in text


def test_partial_overlap_keeps_unique_information_from_both_sources(fake_ocr):
    """M4: ortak kısım tekrar edilmez, her iki kaynağın benzersiz bilgisi korunur."""
    fake_ocr("FEN ISLERI MUDURLUGUNE, Mahallemizde bulunan yol cukurunun onarilmasini talep ediyorum.")
    text = extract_text(make_scan_pdf("Yol cukuru dilekcesi", layer="Evrak No: 2026/4417 Fen Isleri Mudurlugu"), "pdf")

    assert "2026/4417" in text  # yalnızca gömülü metinde
    assert "cukurunun" in text  # yalnızca OCR metninde
    assert text.lower().count("mudurlug") <= 2  # ortak kısım sınırsız tekrarlanmaz


def test_letter_shaped_broken_layer_is_kept_next_to_ocr_text(fake_ocr):
    """Bilinen sınır: harf görünümlü bozuk katman kelime ürettiği için korunur; OCR metni yine tam."""
    fake_ocr()
    text = extract_text(make_scan_pdf("Taranmis dilekce", layer="qwzxk jvbnm xzqwe rtyup"), "pdf")

    assert SCAN_OCR_TEXT in text
    assert "qwzxk" in text


def test_scanned_page_with_healthy_long_text_layer_is_not_ocred(fake_ocr):
    """Sağlıklı uzun katmanı olan tarama yeniden OCR'lanmaz."""
    pages = fake_ocr()
    layer = "\n".join(["Bu taranmis sayfanin metin katmani saglikli ve yeterince uzundur."] * 4)

    assert len(normalize_text(layer)) > file_service.OCR_SHORT_TEXT_MAX
    extract_text(make_scan_pdf("Taranmis dilekce", layer=layer), "pdf")

    assert pages == []


@pytest.mark.parametrize("coverage, ocr_expected", [(0.45, False), (0.55, True)])
def test_image_coverage_threshold_decides_structural_ocr(fake_ocr, coverage, ocr_expected):
    pages = fake_ocr()

    extract_text(make_scan_pdf("Taranmis dilekce", layer=GARBAGE_LAYER, coverage=coverage), "pdf")

    assert pages == ([0] if ocr_expected else [])
    assert file_service.OCR_COVERAGE_MIN == 0.5


@pytest.mark.parametrize("extra, ocr_expected", [(0, True), (1, False)])
def test_short_text_threshold_decides_structural_ocr(fake_ocr, extra, ocr_expected):
    pages = fake_ocr()
    content = make_scan_pdf("Taranmis dilekce", layer=layer_of_length(file_service.OCR_SHORT_TEXT_MAX + extra))

    with pymupdf.open(stream=content, filetype="pdf") as pdf:  # eşiğin iki yanında tam uzunluk
        assert len(normalize_text(pdf[0].get_text())) == file_service.OCR_SHORT_TEXT_MAX + extra
    extract_text(content, "pdf")

    assert pages == ([0] if ocr_expected else [])
    assert file_service.OCR_SHORT_TEXT_MAX == 200


def test_digital_page_with_background_image_keeps_its_text(fake_ocr):
    """Bilinen yanlış pozitif (filigran/arka plan): OCR tetiklense de gömülü bilgi kaybolmaz."""
    fake_ocr("ORNEKTIR Sayi: 89-3345 Tarih: 18.09.2026")
    text = extract_text(make_scan_pdf("ORNEKTIR", layer="Sayi: 89-3345 Tarih: 18.09.2026"), "pdf")

    assert "89-3345" in text and "18.09.2026" in text


@pytest.mark.parametrize(
    "page_text",
    [
        "Evrak Kayit No: 2026/4417 Tarih: 18.09.2026",  # tarih
        "T.C. Kimlik No: 12345678901 VKN: 1234567890",  # TCKN / VKN
        "2026/1 1.250,00 2026/2 980,50 2026/3 12.400,00",  # sayı tablosu
        "1. 2. 3. 4. 5. 6. 7. 8. 9. 10.",  # madde numaraları
        "Sayfa 1 / 12 - 2026/4417",  # sayfa numarası
        "SELECT * FROM documents WHERE id = 1;",  # kod
    ],
)
def test_short_digital_pages_without_images_are_never_ocred(fake_ocr, page_text):
    """Görüntüsü olmayan kısa dijital sayfalar içeriğine bakılmaksızın OCR'lanmaz."""
    pages = fake_ocr()

    assert extract_text(make_pdf(page_text), "pdf") == normalize_text(page_text)
    assert pages == []


def test_structural_ocr_failure_falls_back_to_embedded_text(monkeypatch, caplog):
    """OCR hatası yeni hata sınıfı üretmez; sayfa gömülü metniyle devam eder, teknik detay sızmaz."""
    secret = "TARANMIS_SAYFA_ICERIGI"

    def boom(*args, **kwargs):
        raise RuntimeError(f"tesseract patladi: {secret}")

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", "/tessdata")
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", boom)

    with caplog.at_level(logging.WARNING):
        text = extract_text(make_scan_pdf(secret, layer="Evrak No: 2026/4417"), "pdf")

    assert text == "Evrak No: 2026/4417"
    assert secret not in caplog.text


def test_document_order_is_kept_across_structural_ocr_pages(fake_ocr):
    """Çok sayfalı belgede sayfa sırası korunur."""
    fake_ocr()
    with pymupdf.open() as pdf:
        for part in (make_pdf("Birinci sayfanin yeterli uzunlukta metni."),
                     make_scan_pdf("Taranmis sayfa", layer=GARBAGE_LAYER),
                     make_pdf("Ucuncu sayfanin yeterli uzunlukta metni.")):
            with pymupdf.open(stream=part, filetype="pdf") as source:
                pdf.insert_pdf(source)
        content = pdf.tobytes()

    text = extract_text(content, "pdf")

    assert text.index("Birinci") < text.index("KADIKOY") < text.index("Ucuncu")


def test_ocr_dpi_is_the_same_on_hybrid_and_structural_pages(monkeypatch):
    """P1 hybrid ve P2 yapısal yolda da OCR yalnız gereken sayfada ve aynı çözünürlükle çağrılır."""
    calls = []

    def fake_ocr(self, *args, **kwargs):
        calls.append((self.number, kwargs["language"], kwargs["dpi"]))
        return self.get_textpage()

    monkeypatch.setattr(file_service.settings, "TESSDATA_PREFIX", "/tessdata")
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", fake_ocr)

    # P1: kapak metin sayfası + taranmış sayfa -> yalnız 2. sayfa
    extract_text(make_mixed_pdf(("text", "EVRAK KAYIT FORMU Bu belge kayit sistemine alinmistir."),
                                ("image", "Asil dilekce taranmis sayfada")), "pdf")
    assert calls == [(1, "tur", 400)]

    # P2: bozuk metin katmanı taşıyan taranmış sayfa -> o sayfa
    calls.clear()
    extract_text(make_scan_pdf("Taranmis dilekce metni", layer=GARBAGE_LAYER), "pdf")
    assert calls == [(0, "tur", 400)]
